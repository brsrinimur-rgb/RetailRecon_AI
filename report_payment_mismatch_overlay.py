from __future__ import annotations

import io
import re
from collections import defaultdict

from openpyxl import load_workbook

import report_export


def _norm_ref(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper().strip())


def _to_float(value):
    try:
        if value is None or value == "":
            return None
        return float(value)
    except Exception:
        return None


def _header_map(ws, header_row=4):
    return {
        str(ws.cell(header_row, c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
    }


def _find_payment_mismatch_pairs(ws, header_row=4, amount_tolerance=0.005):
    """
    Read-only detection over the already-produced transaction report.

    Only pairs two EXISTING exception rows when:
      - normalized Auth Code is the same,
      - one row is D365-side only and one row is POS/provider-side only,
      - exact amount is the same within tolerance,
      - D365 Tender and POS Tender are both present and different,
      - the candidate pair is unique on both sides.

    It never creates/deletes rows and never changes a Matched row.
    """
    h = _header_map(ws, header_row)
    required = [
        "Auth Code", "D365 Total", "POS Total",
        "D365 Tender", "POS Tender", "Status", "Remarks"
    ]
    if any(name not in h for name in required):
        return []

    candidates_d365 = []
    candidates_pos = []

    for r in range(header_row + 1, ws.max_row + 1):
        status = str(ws.cell(r, h["Status"]).value or "").strip()
        if status == "Matched":
            continue

        ref = _norm_ref(ws.cell(r, h["Auth Code"]).value)
        d_amt = _to_float(ws.cell(r, h["D365 Total"]).value)
        p_amt = _to_float(ws.cell(r, h["POS Total"]).value)
        d_pay = str(ws.cell(r, h["D365 Tender"]).value or "").upper().strip()
        p_pay = str(ws.cell(r, h["POS Tender"]).value or "").upper().strip()

        if not ref:
            continue

        d_amt = 0.0 if d_amt is None else d_amt
        p_amt = 0.0 if p_amt is None else p_amt

        if d_amt > 0 and abs(p_amt) <= amount_tolerance and d_pay:
            candidates_d365.append((r, ref, d_amt, d_pay))
        elif p_amt > 0 and abs(d_amt) <= amount_tolerance and p_pay:
            candidates_pos.append((r, ref, p_amt, p_pay))

    # Build candidate relations first, then enforce one-to-one uniqueness.
    d_to_p = defaultdict(list)
    p_to_d = defaultdict(list)

    for dr, dref, damt, dpay in candidates_d365:
        for pr, pref, pamt, ppay in candidates_pos:
            if dref != pref:
                continue
            if abs(damt - pamt) > amount_tolerance:
                continue
            if dpay == ppay:
                continue
            d_to_p[dr].append((pr, ppay, pamt))
            p_to_d[pr].append((dr, dpay, damt))

    pairs = []
    for dr, plist in d_to_p.items():
        if len(plist) != 1:
            continue
        pr, ppay, pamt = plist[0]
        if len(p_to_d.get(pr, [])) != 1:
            continue

        drow = next(x for x in candidates_d365 if x[0] == dr)
        _, ref, damt, dpay = drow
        pairs.append((dr, pr, ref, damt, dpay, ppay))

    return pairs


def _apply_overlay_to_sheet(ws, header_row=4, amount_tolerance=0.005):
    h = _header_map(ws, header_row)
    if "Status" not in h or "Remarks" not in h:
        return 0

    pairs = _find_payment_mismatch_pairs(
        ws, header_row=header_row, amount_tolerance=amount_tolerance
    )

    changed = 0
    for d365_row, pos_row, ref, amount, d365_pay, pos_pay in pairs:
        remark = (
            f"Auth Code and amount matched (SAR {amount:,.2f}), but payment method differs: "
            f"D365 = {d365_pay}, Provider = {pos_pay}. "
            "Please verify/correct the payment method in D365 Store Tender."
        )
        for row in (d365_row, pos_row):
            # Safety: never touch an already Matched row.
            current = str(ws.cell(row, h["Status"]).value or "").strip()
            if current == "Matched":
                continue
            ws.cell(row, h["Status"]).value = "Payment Method Mismatch"
            ws.cell(row, h["Remarks"]).value = remark
            changed += 1
    return changed


def create_reconciliation_pack(result, tolerance=1.0):
    """
    EXPORT OVERLAY ONLY.

    1. Calls the existing report_export.create_reconciliation_pack unchanged.
    2. Re-opens the produced workbook in memory.
    3. Re-labels only proven payment-method mismatch exception pairs.
    4. Leaves Dashboard counts, matched rows, matching engine, mappings and
       normalization untouched.
    """
    original_bytes = report_export.create_reconciliation_pack(result, tolerance)

    wb = load_workbook(io.BytesIO(original_bytes))

    # Update the two user-facing transaction/exception sheets only.
    for sheet_name in ("Transaction_Reconciliation", "Exceptions"):
        if sheet_name in wb.sheetnames:
            _apply_overlay_to_sheet(
                wb[sheet_name],
                header_row=4,
                amount_tolerance=0.005,
            )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
