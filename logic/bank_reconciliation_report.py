"""
logic/bank_reconciliation_report.py

Finance-facing Bank Reconciliation Control Tower report (page 38).

This module does NOT re-parse bank statements or re-run settlement matching.
It reads the settlement batch results the Settlement Batch Engine (page 18)
already computes -- st.session_state["ct_result"]["settlement_batches"] /
["settlement_bank_unmatched"] -- built by:
    core.build_card_settlement_batches()            (MADA/VISA/MASTERCARD/AMEX)
    core.normalize_tabby_payout / tamara_payout / tap_payout()
    logic.bank_settlement_extension.reconcile_card_batches_advanced()
    logic.bank_settlement_extension.reconcile_provider_batches_to_rajhi()

Every one of those functions already decides Settlement Status ("BANK
RECEIVED" / "BANK RECEIPT PENDING" / "BANK REVIEW REQUIRED", or any other
status text a future matcher adds -- this module never hardcodes an
exhaustive status list, it always falls through to surfacing whatever text
is actually there) and Bank Match Rule (the audit trail of which evidence
tied a settlement to a bank credit). This module's job is purely to turn
that already-decided data into a finance-readable Dashboard / Provider
Summary / Store Summary / Aging / Exceptions / Control Totals report and a
polished multi-sheet Excel workbook -- nothing here changes which
transactions matched or why. Page 11 (Bank Settlement Audit) remains the
detailed, transaction-level settlement audit; this page is the management /
month-end CONSOLIDATED view built on top of page 18's batch-level output --
the two are not duplicates of each other.

Known, stated limitation carried over from the underlying engine: matching
is strictly 1 settlement batch <-> 1 bank credit (see
reconcile_card_batches_advanced / reconcile_provider_batches_to_rajhi).
Many-to-many settlement-to-bank-credit matching (one bank credit funding
several settlements, or one settlement paid across several credits) is NOT
implemented anywhere in this codebase yet. Where that happens in real data,
today's engine correctly leaves those rows unresolved (BANK RECEIPT PENDING
on the settlement side, an unmatched credit on the bank side) rather than
guessing -- consistent with this project's "never guess" matching
discipline. This report surfaces that honestly rather than papering over it;
extending the underlying matcher to support many-to-many is a separate,
larger change, deliberately NOT made here (V69 only touches this reporting
module, per an explicit decision to freeze the matching engine while this
layer is proven against real data first).

V69 changes (on top of V68), from a real code review:
  - HIGH-priority thresholds are now inclusive (>=) so an amount/age exactly
    at the configured threshold counts as HIGH, not MEDIUM.
  - A settlement whose Settlement Date is in the future is no longer folded
    into ordinary aging (which could previously go negative) -- it's now its
    own "Future Settlement Date - Data Validation Required" exception.
  - Dashboard and Provider Summary now report BOTH a Group Match % (count of
    clean settlement groups / total groups) and a Value Match % (SAR value
    of clean groups / total SAR value) -- a portfolio can look highly
    matched by count while a single large exception dominates the SAR
    exposure, and Value Match % is what catches that.
  - A bank-side control identity is now printed: Total Bank Credits in Scope
    = Allocated Bank Credits (landed against a settlement) + Unidentified
    Bank Credits (landed but unmatched). This is a partition of the same
    data this module already has, not an independently-sourced bank total,
    and is presented as such.
  - Provider Summary gained Group Match %, Value Match %, Outstanding,
    Average Delay, Oldest Outstanding per provider.
  - New Store x Provider x Payment Type summary sheet, so a store's
    outstanding balance can be attributed to a specific provider/tender.
  - Exception priority is now a deterministic 4-tier rule (CRITICAL / HIGH /
    MEDIUM / LOW) instead of HIGH/MEDIUM only, including a new, purely
    rule-based "possible duplicate settlement batch" integrity check
    (same Provider + Store + Payment Type + Settlement Date + Expected
    Bank Amount appearing on more than one settlement batch) flagged
    CRITICAL regardless of amount.
  - The Excel export now has frozen header rows, autofilters, SAR/date/
    percent number formats, sized columns, and priority/status colour
    fills, instead of plain unformatted sheets.
"""
from __future__ import annotations

import numpy as np
import pandas as pd

AGING_BINS = [-1, 2, 5, 10, 10_000_000]
AGING_LABELS = ["0-2 days", "3-5 days", "6-10 days", ">10 days"]

RECEIVED = "BANK RECEIVED"
PENDING = "BANK RECEIPT PENDING"
REVIEW = "BANK REVIEW REQUIRED"

DUPLICATE_LABEL = "Possible Duplicate Settlement Batch"
FUTURE_DATE_LABEL = "Future Settlement Date - Data Validation Required"

DUPLICATE_KEY = ["Provider", "Store Code", "Payment Type", "Settlement Date", "Expected Bank Amount"]

_NUMERIC_COLS = [
    "Gross Amount", "Refund Amount", "Fee Amount", "VAT Amount",
    "Expected Bank Amount", "Actual Bank Amount", "Bank Difference",
    "Transaction Count",
]


def _num(s) -> pd.Series:
    return pd.to_numeric(s, errors="coerce")


def build_detail(batches: pd.DataFrame) -> pd.DataFrame:
    """
    Normalize settlement_batches (from page 18 / core.py) into the detail
    report, adding Delay Days (received rows), Age Days / Age Bucket
    (outstanding rows, excluding future-dated ones -- see Future Dated
    below), and a Possible Duplicate flag. Safe against missing optional
    columns -- card batches and provider/BNPL batches don't carry exactly
    the same column set (e.g. ANB Commission/VAT only exist for card
    batches), and this function is the one place that reconciles the two
    shapes into one consistent report.
    """
    if batches is None or batches.empty:
        return pd.DataFrame()

    x = batches.copy()

    for c in _NUMERIC_COLS:
        x[c] = _num(x[c]) if c in x.columns else np.nan

    for c in ["Provider", "Store Code", "Payment Type", "Settlement Status",
              "Bank Match Rule", "Settlement Review Reason", "Bank Reference",
              "Settlement Batch ID", "Merchant", "Terminal ID"]:
        if c not in x.columns:
            x[c] = ""
        x[c] = x[c].fillna("").astype(str)

    x["Settlement Date"] = pd.to_datetime(x.get("Settlement Date"), errors="coerce")
    x["Bank Date"] = pd.to_datetime(x.get("Bank Date"), errors="coerce") if "Bank Date" in x.columns else pd.NaT

    received = x["Settlement Status"].eq(RECEIVED)
    today = pd.Timestamp.today().normalize()

    x["Delay Days"] = np.where(
        received & x["Settlement Date"].notna() & x["Bank Date"].notna(),
        (x["Bank Date"] - x["Settlement Date"]).dt.days,
        np.nan,
    )

    future_dated = (~received) & x["Settlement Date"].notna() & (x["Settlement Date"] > today)
    x["Future Dated"] = future_dated

    raw_age = (today - x["Settlement Date"]).dt.days
    x["Age Days"] = np.where((~received) & x["Settlement Date"].notna() & (~future_dated), raw_age, np.nan)
    x["Days Until Settlement"] = np.where(future_dated, -raw_age, np.nan)

    age_for_bucket = pd.to_numeric(x["Age Days"], errors="coerce")
    x["Age Bucket"] = pd.cut(age_for_bucket, bins=AGING_BINS, labels=AGING_LABELS)

    # Possible duplicate settlement batch: same natural key (provider, store,
    # payment type, settlement date, expected amount) appearing on more than
    # one distinct settlement batch. A purely deterministic integrity check
    # over data this module already has -- it does not touch or second-guess
    # the underlying matcher's own identity/dedup rules.
    valid_key = x[DUPLICATE_KEY].notna().all(axis=1)
    x["_DupGroupSize"] = 1
    if valid_key.any():
        sizes = x.loc[valid_key].groupby(DUPLICATE_KEY)["Settlement Batch ID"].transform("size")
        x.loc[valid_key, "_DupGroupSize"] = sizes
    x["Possible Duplicate"] = x["_DupGroupSize"] > 1
    x = x.drop(columns=["_DupGroupSize"])

    return x


def classify_exceptions(
    detail: pd.DataFrame,
    amount_threshold: float = 5000.0,
    age_high_days: int = 5,
    delay_high_days: int = 5,
    low_amount_threshold: float = 500.0,
) -> pd.DataFrame:
    """
    Every row that isn't a clean BANK RECEIVED-with-zero-difference is an
    exception -- plus, regardless of status, any row flagged Possible
    Duplicate or Future Dated. The exception label mirrors the underlying
    Settlement Status whenever this module doesn't have a more specific
    finance label for it, so a status this report has never seen before
    still shows up instead of silently vanishing.

    Priority is a deterministic 4-tier rule, not a hidden scoring model:
      CRITICAL - possible duplicate settlement batch (control risk,
                 regardless of amount).
      HIGH     - review-required (multiple bank-credit candidates),
                 a future-dated settlement (data validation issue), or the
                 expected amount / age / delay meets or exceeds the
                 configured threshold (>=, so a value exactly at the
                 threshold counts as HIGH).
      MEDIUM   - amount differences and late settlements that don't clear
                 the HIGH bar, or anything of moderate size/age.
      LOW      - small, fresh items kept only for monitoring.
    All thresholds are parameters so Finance can tune them without a code
    change once this is wired to a settings page, if wanted.
    """
    if detail is None or detail.empty:
        return pd.DataFrame()

    x = detail.copy()

    def _label(r):
        if r.get("Possible Duplicate", False):
            return DUPLICATE_LABEL
        if r.get("Future Dated", False):
            return FUTURE_DATE_LABEL
        status = r.get("Settlement Status", "")
        if status == RECEIVED:
            diff = r.get("Bank Difference", np.nan)
            if pd.notna(diff) and abs(diff) > 0.01:
                return "Amount Difference"
            delay = r.get("Delay Days", np.nan)
            if pd.notna(delay) and delay >= delay_high_days:
                return "Late Settlement"
            return ""
        if status == PENDING:
            return "Settlement Not Received"
        if status == REVIEW:
            return "Review - Multiple Candidates"
        if status:
            return status
        return "Unclassified"

    x["Exception"] = x.apply(_label, axis=1)
    exc = x[x["Exception"] != ""].copy()
    if exc.empty:
        return exc

    def _priority(r):
        label = r["Exception"]
        if label == DUPLICATE_LABEL:
            return "CRITICAL"
        if label in {"Review - Multiple Candidates", FUTURE_DATE_LABEL}:
            return "HIGH"

        amt = abs(r.get("Expected Bank Amount", 0) or 0)
        age = r.get("Age Days", np.nan)
        delay = r.get("Delay Days", np.nan)
        worst = age if pd.notna(age) else (delay if pd.notna(delay) else None)

        if amt >= amount_threshold or (worst is not None and worst >= age_high_days):
            return "HIGH"
        if label in {"Amount Difference", "Late Settlement"}:
            return "MEDIUM"
        if amt < low_amount_threshold and (worst is None or worst <= 1):
            return "LOW"
        return "MEDIUM"

    exc["Priority"] = exc.apply(_priority, axis=1)

    action_map = {
        "Settlement Not Received": "Claim Provider / Bank",
        "Amount Difference": "Investigate",
        "Late Settlement": "Monitor",
        "Review - Multiple Candidates": "Investigate",
        DUPLICATE_LABEL: "Investigate Urgently",
        FUTURE_DATE_LABEL: "Fix Source Data",
        "Unclassified": "Review",
    }
    exc["Action"] = exc["Exception"].map(action_map).fillna("Review")

    def _age(r):
        a = r.get("Age Days", np.nan)
        if pd.notna(a):
            return f"{int(a)} days"
        d = r.get("Delay Days", np.nan)
        if pd.notna(d):
            return f"{int(d)} days"
        u = r.get("Days Until Settlement", np.nan)
        return f"in {int(u)} days" if pd.notna(u) else ""

    exc["Age"] = exc.apply(_age, axis=1)
    return exc


def dashboard_kpis(detail: pd.DataFrame, bank_unmatched: pd.DataFrame | None = None) -> dict:
    if detail is None or detail.empty:
        return {
            "Expected Settlements": 0.0, "Received in Bank": 0.0, "Matched": 0.0,
            "Amount Differences": 0.0, "Not Yet Received": 0.0, "Review Required": 0.0,
            "Allocated Bank Credits": 0.0, "Unidentified Bank Credits": 0.0,
            "Total Bank Credits in Scope": 0.0,
            "Total Settlement Groups": 0, "Fully Matched": 0, "Exceptions": 0,
            "Possible Duplicates": 0, "Future Dated Settlements": 0,
            "Group Match %": 0.0, "Value Match %": 0.0,
            "Average Settlement Delay (days)": None, "Oldest Outstanding (days)": None,
        }

    x = detail
    received = x["Settlement Status"].eq(RECEIVED)
    pending = x["Settlement Status"].eq(PENDING)
    review = x["Settlement Status"].eq(REVIEW)
    clean = received & (x["Bank Difference"].abs() <= 0.01)
    diffed = received & (x["Bank Difference"].abs() > 0.01)
    dup = x.get("Possible Duplicate", pd.Series(False, index=x.index)).fillna(False)
    future = x.get("Future Dated", pd.Series(False, index=x.index)).fillna(False)

    allocated_amt = float(x.loc[received, "Actual Bank Amount"].sum())

    unmatched_amt = 0.0
    if bank_unmatched is not None and not bank_unmatched.empty:
        col = next((c for c in ["Credit", "Bank Amount", "Amount"] if c in bank_unmatched.columns), None)
        if col:
            unmatched_amt = float(_num(bank_unmatched[col]).fillna(0).sum())

    total = len(x)
    fully_matched = int(clean.sum())
    expected_total = float(x["Expected Bank Amount"].sum())
    clean_expected = float(x.loc[clean, "Expected Bank Amount"].sum())

    return {
        "Expected Settlements": expected_total,
        "Received in Bank": allocated_amt,
        "Matched": float(x.loc[clean, "Actual Bank Amount"].sum()),
        "Amount Differences": float(x.loc[diffed, "Bank Difference"].abs().sum()),
        "Not Yet Received": float(x.loc[pending, "Expected Bank Amount"].sum()),
        "Review Required": float(x.loc[review, "Expected Bank Amount"].sum()),
        "Allocated Bank Credits": allocated_amt,
        "Unidentified Bank Credits": unmatched_amt,
        "Total Bank Credits in Scope": round(allocated_amt + unmatched_amt, 2),
        "Total Settlement Groups": total,
        "Fully Matched": fully_matched,
        "Exceptions": total - fully_matched,
        "Possible Duplicates": int(dup.sum()),
        "Future Dated Settlements": int(future.sum()),
        "Group Match %": round(fully_matched / total * 100, 2) if total else 0.0,
        "Value Match %": round(clean_expected / expected_total * 100, 2) if expected_total else 0.0,
        "Average Settlement Delay (days)": (
            round(float(x.loc[received, "Delay Days"].mean()), 1)
            if received.any() and x.loc[received, "Delay Days"].notna().any() else None
        ),
        "Oldest Outstanding (days)": (
            int(x.loc[~received, "Age Days"].max())
            if (~received).any() and x.loc[~received, "Age Days"].notna().any() else None
        ),
    }


def provider_summary(detail: pd.DataFrame) -> pd.DataFrame:
    if detail is None or detail.empty:
        return pd.DataFrame()
    x = detail.copy()
    x["_clean"] = x["Settlement Status"].eq(RECEIVED) & (x["Bank Difference"].abs() <= 0.01)
    x["_received"] = x["Settlement Status"].eq(RECEIVED)

    rows = []
    for provider, g in x.groupby("Provider", dropna=False):
        expected = float(g["Expected Bank Amount"].sum())
        received_amt = float(g.loc[g["_received"], "Actual Bank Amount"].sum())
        clean_expected = float(g.loc[g["_clean"], "Expected Bank Amount"].sum())
        groups = len(g)
        matched = int(g["_clean"].sum())
        diff = round(expected - received_amt, 2)
        rows.append({
            "Provider": provider,
            "Expected": expected,
            "Bank Received": received_amt,
            "Difference": diff,
            "Outstanding": diff,
            "Groups": groups,
            "Matched": matched,
            "Exceptions": groups - matched,
            "Group Match %": round(matched / groups * 100, 2) if groups else 0.0,
            "Value Match %": round(clean_expected / expected * 100, 2) if expected else 0.0,
            "Average Delay (days)": (
                round(float(g.loc[g["_received"], "Delay Days"].mean()), 1)
                if g["_received"].any() and g.loc[g["_received"], "Delay Days"].notna().any() else None
            ),
            "Oldest Outstanding (days)": (
                int(g.loc[~g["_received"], "Age Days"].max())
                if (~g["_received"]).any() and g.loc[~g["_received"], "Age Days"].notna().any() else None
            ),
        })
    return pd.DataFrame(rows)


def store_summary(detail: pd.DataFrame) -> pd.DataFrame:
    if detail is None or detail.empty:
        return pd.DataFrame()
    x = detail.copy()
    x["_clean"] = x["Settlement Status"].eq(RECEIVED) & (x["Bank Difference"].abs() <= 0.01)
    g = x.groupby("Store Code", dropna=False).agg(
        Expected=("Expected Bank Amount", "sum"),
        Received=("Actual Bank Amount", "sum"),
        Groups=("Settlement Batch ID", "count"),
        Matched=("_clean", "sum"),
    ).reset_index()
    g["Outstanding"] = (g["Expected"] - g["Received"]).round(2)
    g["Exceptions"] = g["Groups"] - g["Matched"]
    g["Group Match %"] = np.where(g["Groups"] > 0, (g["Matched"] / g["Groups"] * 100).round(2), 0.0)
    return g[["Store Code", "Expected", "Received", "Outstanding", "Groups", "Matched", "Exceptions", "Group Match %"]]


def store_provider_summary(detail: pd.DataFrame) -> pd.DataFrame:
    """Store x Provider x Payment Type breakdown -- answers "store 601 has
    SAR 20,000 outstanding, but which provider/tender is it?" which the
    store-only summary above can't."""
    if detail is None or detail.empty:
        return pd.DataFrame()
    x = detail.copy()
    x["_clean"] = x["Settlement Status"].eq(RECEIVED) & (x["Bank Difference"].abs() <= 0.01)
    g = x.groupby(["Store Code", "Provider", "Payment Type"], dropna=False).agg(
        Expected=("Expected Bank Amount", "sum"),
        Received=("Actual Bank Amount", "sum"),
        Groups=("Settlement Batch ID", "count"),
        Matched=("_clean", "sum"),
    ).reset_index()
    g["Outstanding"] = (g["Expected"] - g["Received"]).round(2)
    g["Exceptions"] = g["Groups"] - g["Matched"]
    g["Group Match %"] = np.where(g["Groups"] > 0, (g["Matched"] / g["Groups"] * 100).round(2), 0.0)
    return g.sort_values(["Store Code", "Provider", "Payment Type"]).reset_index(drop=True)


def daily_summary(detail: pd.DataFrame) -> pd.DataFrame:
    if detail is None or detail.empty:
        return pd.DataFrame()
    x = detail.copy()
    x["Settlement Day"] = x["Settlement Date"].dt.date
    x["_clean"] = x["Settlement Status"].eq(RECEIVED) & (x["Bank Difference"].abs() <= 0.01)
    g = x.groupby("Settlement Day", dropna=False).agg(
        Expected=("Expected Bank Amount", "sum"),
        Received=("Actual Bank Amount", "sum"),
        Groups=("Settlement Batch ID", "count"),
        Matched=("_clean", "sum"),
    ).reset_index()
    return g.sort_values("Settlement Day")


def matching_audit(detail: pd.DataFrame) -> pd.DataFrame:
    """Breakdown by the literal Bank Match Rule / Settlement Status text --
    this IS the audit trail of which evidence level (Settlement ID + exact
    amount, provider + date-window + amount, etc.) tied each group, using
    whatever rule text the underlying matcher already produced."""
    if detail is None or detail.empty:
        return pd.DataFrame()
    x = detail.copy()
    x["Match Basis"] = np.where(
        x["Bank Match Rule"].astype(str).str.strip() != "",
        x["Bank Match Rule"],
        x["Settlement Status"],
    )
    g = x.groupby(["Settlement Status", "Match Basis"], dropna=False).agg(
        Groups=("Settlement Batch ID", "count"),
        Amount=("Expected Bank Amount", "sum"),
    ).reset_index().sort_values(["Settlement Status", "Groups"], ascending=[True, False])
    return g


def aging_summary(detail: pd.DataFrame) -> pd.DataFrame:
    """Outstanding (non-received) settlements, excluding future-dated ones
    (those are their own exception/control item, not an aging concern)."""
    if detail is None or detail.empty:
        return pd.DataFrame()
    future = detail.get("Future Dated", pd.Series(False, index=detail.index)).fillna(False)
    x = detail[detail["Settlement Status"].ne(RECEIVED) & ~future].copy()
    if x.empty:
        return pd.DataFrame()
    g = x.groupby("Age Bucket", dropna=False, observed=False).agg(
        Groups=("Settlement Batch ID", "count"),
        Amount=("Expected Bank Amount", "sum"),
    ).reset_index()
    g["Age Bucket"] = g["Age Bucket"].astype(str)
    order = {lbl: i for i, lbl in enumerate(AGING_LABELS)}
    g["_ord"] = g["Age Bucket"].map(order).fillna(99)
    return g.sort_values("_ord").drop(columns="_ord")


def control_totals(detail: pd.DataFrame, bank_unmatched: pd.DataFrame | None) -> pd.DataFrame:
    """
    Proves the detail sheet ties back to the dashboard, two ways:

    1. Settlement side: every row's Expected Bank Amount is bucketed into
       exactly one of clean-matched / received-with-difference / pending /
       review / other-status, and the sum is shown tying to the total.
    2. Bank side: Total Bank Credits in Scope = Allocated Bank Credits
       (landed against a settlement, whether clean or diffed) + Unidentified
       Bank Credits (landed but unmatched to anything).

    Because every number here is derived from the same `detail` frame (and
    the bank_unmatched frame this module already has), both identities tie
    by construction -- this is a printed proof for the reader and a guard
    against a future edit to this module accidentally breaking the
    identity, not an independent check of the underlying matching engine's
    decisions or of the bank statement's true total credits.
    """
    if detail is None or detail.empty:
        return pd.DataFrame()

    x = detail
    received = x["Settlement Status"].eq(RECEIVED)
    clean = received & (x["Bank Difference"].abs() <= 0.01)
    diffed = received & (x["Bank Difference"].abs() > 0.01)
    pending = x["Settlement Status"].eq(PENDING)
    review = x["Settlement Status"].eq(REVIEW)
    other = ~(clean | diffed | pending | review)

    expected_total = float(x["Expected Bank Amount"].sum())
    clean_exp = float(x.loc[clean, "Expected Bank Amount"].sum())
    diffed_exp = float(x.loc[diffed, "Expected Bank Amount"].sum())
    pending_exp = float(x.loc[pending, "Expected Bank Amount"].sum())
    review_exp = float(x.loc[review, "Expected Bank Amount"].sum())
    other_exp = float(x.loc[other, "Expected Bank Amount"].sum())
    accounted = clean_exp + diffed_exp + pending_exp + review_exp + other_exp

    allocated_amt = float(x.loc[received, "Actual Bank Amount"].sum())
    unmatched_amt = 0.0
    if bank_unmatched is not None and not bank_unmatched.empty:
        col = next((c for c in ["Credit", "Bank Amount", "Amount"] if c in bank_unmatched.columns), None)
        if col:
            unmatched_amt = float(_num(bank_unmatched[col]).fillna(0).sum())
    bank_side_total = allocated_amt + unmatched_amt

    dup_count = int(x.get("Possible Duplicate", pd.Series(False, index=x.index)).fillna(False).sum())
    future_count = int(x.get("Future Dated", pd.Series(False, index=x.index)).fillna(False).sum())

    rows = [
        {"Line": "SETTLEMENT SIDE", "Amount": None},
        {"Line": "Total Expected Settlements", "Amount": round(expected_total, 2)},
        {"Line": "  Matched (clean, zero difference)", "Amount": round(clean_exp, 2)},
        {"Line": "  Received with Amount Difference (expected value of these groups)", "Amount": round(diffed_exp, 2)},
        {"Line": "  Not Yet Received (Pending)", "Amount": round(pending_exp, 2)},
        {"Line": "  Review Required", "Amount": round(review_exp, 2)},
    ]
    if other.any():
        rows.append({"Line": "  Other / Unrecognized Status", "Amount": round(other_exp, 2)})
    rows += [
        {"Line": "Sum of the lines above", "Amount": round(accounted, 2)},
        {"Line": "Ties to Total Expected (should be 0.00)", "Amount": round(expected_total - accounted, 2)},
        {"Line": "", "Amount": None},
        {"Line": "BANK SIDE", "Amount": None},
        {"Line": "Allocated Bank Credits (landed against a settlement)", "Amount": round(allocated_amt, 2)},
        {"Line": "Unidentified Bank Credits (landed, not tied to any settlement)", "Amount": round(unmatched_amt, 2)},
        {"Line": "Total Bank Credits in Scope", "Amount": round(bank_side_total, 2)},
        {"Line": "Ties (Allocated + Unidentified - Total, should be 0.00)",
         "Amount": round(allocated_amt + unmatched_amt - bank_side_total, 2)},
        {"Line": "", "Amount": None},
        {"Line": "DATA QUALITY", "Amount": None},
        {"Line": "Possible Duplicate Settlement Batches (count)", "Amount": dup_count},
        {"Line": "Future-Dated Settlements (count)", "Amount": future_count},
    ]
    return pd.DataFrame(rows)


def build_report(batches: pd.DataFrame, bank_unmatched: pd.DataFrame | None = None,
                  amount_threshold: float = 5000.0, age_high_days: int = 5,
                  delay_high_days: int = 5, low_amount_threshold: float = 500.0) -> dict:
    """Single entry point: returns every sheet/table this report produces."""
    detail = build_detail(batches)
    return {
        "detail": detail,
        "exceptions": classify_exceptions(detail, amount_threshold, age_high_days, delay_high_days, low_amount_threshold),
        "kpis": dashboard_kpis(detail, bank_unmatched),
        "provider_summary": provider_summary(detail),
        "store_summary": store_summary(detail),
        "store_provider_summary": store_provider_summary(detail),
        "daily_summary": daily_summary(detail),
        "aging": aging_summary(detail),
        "matching_audit": matching_audit(detail),
        "control_totals": control_totals(detail, bank_unmatched),
        "bank_unmatched": bank_unmatched if bank_unmatched is not None else pd.DataFrame(),
    }


DETAIL_DISPLAY_COLS = [
    ("Bank Date", "Bank Date"),
    ("Provider", "Provider"),
    ("Store Code", "Store"),
    ("Settlement Batch ID", "Settlement ID"),
    ("Settlement Date", "Sales/Settlement Date"),
    ("Gross Amount", "Gross SAR"),
    ("Fee Amount", "Fee SAR"),
    ("VAT Amount", "VAT SAR"),
    ("Expected Bank Amount", "Expected Bank SAR"),
    ("Actual Bank Amount", "Actual Bank SAR"),
    ("Bank Difference", "Difference"),
    ("Delay Days", "Delay Days"),
    ("Settlement Status", "Status"),
]

CURRENCY_HEADERS = {
    "Gross SAR", "Fee SAR", "VAT SAR", "Expected Bank SAR", "Actual Bank SAR", "Difference",
    "Expected", "Bank Received", "Received", "Outstanding", "Amount", "Result", "Bank Amount",
    "Credit", "Gross Amount", "Fee Amount", "VAT Amount", "Expected Bank Amount", "Actual Bank Amount",
    "Bank Difference",
}
DATE_HEADERS = {"Bank Date", "Sales/Settlement Date", "Settlement Date", "Settlement Day"}
PERCENT_HEADERS = {"Group Match %", "Value Match %", "Match %"}


def display_detail(detail: pd.DataFrame) -> pd.DataFrame:
    """Renames/orders the detail frame to the exact finance-facing column
    headers, keeping the remaining audit-trail columns (Bank Match Rule,
    Terminal ID, Merchant, Bank Reference, Settlement Review Reason,
    Transaction Count, Age Days/Bucket, Possible Duplicate, Future Dated)
    appended after them so nothing is dropped."""
    if detail is None or detail.empty:
        return pd.DataFrame()
    front = [c for c, _ in DETAIL_DISPLAY_COLS if c in detail.columns]
    rest = [c for c in detail.columns if c not in front]
    out = detail[front + rest].rename(columns=dict(DETAIL_DISPLAY_COLS))
    return out


def _style_worksheet(ws, df: pd.DataFrame) -> None:
    """Applies the shared professional formatting to one written sheet:
    bold/filled header, frozen header row, autofilter, sized columns, and
    SAR/date/percent number formats by column header name."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if df is None or df.empty or ws.max_row < 1:
        return

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        header = str(col_name)
        max_len = max([len(header)] + [len(str(v)) for v in df[col_name].astype(str).head(200)])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)

        num_fmt = None
        if header in PERCENT_HEADERS:
            num_fmt = '0.00"%"'
        elif header in DATE_HEADERS:
            num_fmt = "dd-mmm-yy"
        elif header in CURRENCY_HEADERS:
            num_fmt = "#,##0.00"
        if num_fmt:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = num_fmt


def _fill_rows(ws, df: pd.DataFrame, color_col: str, color_map: dict) -> None:
    """Colours entire data rows based on the value in `color_col` (e.g.
    Priority), using color_map {value: hex}. Applied after _style_worksheet
    so the header formatting isn't overwritten."""
    if df is None or df.empty or color_col not in df.columns:
        return
    _fill_rows_by_series(ws, df[color_col].astype(str), color_map, ncols=len(df.columns))


def _fill_rows_by_series(ws, values: pd.Series, color_map: dict, ncols: int) -> None:
    """Colours entire data rows using a color Series aligned 1:1 with the
    written rows (row 2 = values.iloc[0], etc.) -- used when the colour
    decision depends on more than one column (e.g. clean vs diffed both
    share Settlement Status == BANK RECEIVED, only Bank Difference tells
    them apart) so the deciding column doesn't need to be the visible one."""
    from openpyxl.styles import PatternFill

    for row_idx, value in enumerate(values, start=2):
        hexcode = color_map.get(value)
        if not hexcode:
            continue
        fill = PatternFill("solid", fgColor=hexcode)
        for c in range(1, ncols + 1):
            ws.cell(row=row_idx, column=c).fill = fill


def write_excel(report: dict, path_or_buffer) -> None:
    kpis = report["kpis"]
    dash = pd.DataFrame([{"KPI": k, "Result": v} for k, v in kpis.items()])

    detail_disp = display_detail(report["detail"])
    exc_disp = report["exceptions"]

    with pd.ExcelWriter(path_or_buffer, engine="openpyxl") as writer:
        sheets = {
            "Dashboard": dash,
            "Bank Reconciliation": detail_disp,
            "Exceptions": exc_disp,
            "Aging": report["aging"],
            "Provider Summary": report["provider_summary"],
            "Store Summary": report["store_summary"],
            "Store x Provider Summary": report["store_provider_summary"],
            "Daily Summary": report["daily_summary"],
            "Unmatched Bank Credits": report["bank_unmatched"],
            "Matching Audit": report["matching_audit"],
            "Control Totals": report["control_totals"],
        }
        for name, df in sheets.items():
            (df if df is not None else pd.DataFrame()).to_excel(writer, sheet_name=name, index=False)

        for name, df in sheets.items():
            ws = writer.sheets[name]
            _style_worksheet(ws, df)

        _fill_rows(writer.sheets["Exceptions"], exc_disp, "Priority", {
            "CRITICAL": "FFC7CE", "HIGH": "FFD9B3", "MEDIUM": "FFF2CC", "LOW": "E2EFDA",
        })

        raw_detail = report["detail"]
        if not detail_disp.empty and not raw_detail.empty:
            clean = raw_detail["Settlement Status"].eq(RECEIVED) & (raw_detail["Bank Difference"].abs() <= 0.01)
            exception_row = raw_detail["Settlement Batch ID"].isin(
                set(report["exceptions"]["Settlement Batch ID"]) if not report["exceptions"].empty else set()
            )
            row_labels = pd.Series(
                np.where(clean, "CLEAN", np.where(exception_row, "EXCEPTION", "")),
                index=raw_detail.index,
            )
            _fill_rows_by_series(
                writer.sheets["Bank Reconciliation"], row_labels,
                {"CLEAN": "E2EFDA", "EXCEPTION": "FCE4D6"},
                ncols=len(detail_disp.columns),
            )
